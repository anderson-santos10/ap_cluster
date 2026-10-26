"""Pré-validação da planilha aprovada.

Lê `Criticos Ap Cluster.xlsx` e imprime um relatório. Não salva a planilha
e não abre conexão com o banco.
"""

import argparse
import hashlib
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from criticals.preparation import (
    classify_position_group,
    descriptions_for_review,
    excel_serial_calendar,
    missing_stock,
    movement_key,
    po_cache_relation,
    prepare_base_row,
    read_center,
    referenced_base_row,
    repeated_origins,
    text_from_code,
)

DEFAULT_SOURCE = Path(r"C:\Users\Anderson\Documents\PROJETOS\CZC\Criticos Ap Cluster.xlsx")
EXPECTED_SHEETS = ["Base", "Grafico", "SQVI", "LRR", "Est. Parq.", "Din"]
BASE_HEADERS = (
    "Conc",
    "GP",
    "Descrição",
    "Parque",
    "Planta",
    "LRR",
    "Estoque Parque",
    "Estoque Ok/Nok",
    "Quant. PO ",
    "Previsão no Parque ",
    "Mes/Ano",
    "Coluna4",
    "Coluna3",
    "Status",
    "Quant. Entregue",
    "Data entrega ",
    "Obs",
)
STOCK_HEADERS = (
    None,
    "Centro",
    "Nome 1",
    "Depósito",
    "Descrição do depósito",
    "Material",
    "Texto breve material",
    "UM básica",
    "Lote",
    "Utilização livre",
    "Estoque em trânsito",
    "Em contr.qualidade",
    "Bloqueado",
)
SQVI_INDEXES = {
    1: "conc ",
    7: "Tipo",
    9: "Doc.compra",
    10: "Item",
    13: "CFor",
    17: "Cen.",
    19: "Material",
    21: "Qtd.pedido",
    36: "Qtd.confir.",
    39: "DataConf.",
    44: "Data fim da EM",
    52: "Qtd.saída",
    54: "Fornecimento",
    55: "Qtd.remessa",
}


def main():
    parser = argparse.ArgumentParser(description="Pré-valida a planilha sem gravar dados.")
    parser.add_argument("--origem", default=str(DEFAULT_SOURCE))
    args = parser.parse_args()
    _configure_output()
    source = Path(args.origem)
    if source.name != "Criticos Ap Cluster.xlsx":
        raise SystemExit("A origem precisa ser Criticos Ap Cluster.xlsx.")
    if "LRR BR" in source.name:
        raise SystemExit("LRR BR está fora desta pré-validação.")
    before = _file_stamp(source)
    formula_book = load_workbook(source, read_only=True, data_only=False)
    value_book = load_workbook(source, read_only=True, data_only=True)
    try:
        report = build_report(formula_book, value_book)
    finally:
        formula_book.close()
        value_book.close()
    after = _file_stamp(source)
    report["arquivo"] = before
    report["arquivo_depois"] = after
    if before["sha256"] != after["sha256"] or before["bytes"] != after["bytes"]:
        raise SystemExit("A planilha mudou durante a leitura. Nenhum resultado foi aceito.")
    _print_report(report)


def build_report(formula_book, value_book):
    if list(formula_book.sheetnames) != EXPECTED_SHEETS:
        raise SystemExit(f"Abas inesperadas: {formula_book.sheetnames}")
    base_formulas = _rows(formula_book["Base"])
    base_values = _rows(value_book["Base"])
    stock_values = _rows(value_book["Est. Parq."])
    sqvi_values = _rows(value_book["SQVI"])
    lrr_values = _rows(value_book["LRR"])
    _require_headers(base_values[0], BASE_HEADERS, "Base")
    _require_headers(stock_values[0], STOCK_HEADERS, "Est. Parq.")
    for index, header in SQVI_INDEXES.items():
        if sqvi_values[0][index] != header:
            raise SystemExit(f"Cabeçalho SQVI coluna {index + 1} diferente do esperado.")
    if lrr_values[0][1] != "Material" or lrr_values[0][3] != "Área MRP centro":
        raise SystemExit("Cabeçalho da aba LRR diferente do esperado.")
    if lrr_values[0][8] != "Ponto reabastec.":
        raise SystemExit("Coluna de ponto de reabastecimento não encontrada.")
    base = _base_section(base_formulas, base_values)
    stock = _stock_section(stock_values)
    sqvi = _sqvi_section(sqvi_values)
    return {
        "abas": {
            "Base": len(base["itens"]),
            "SQVI": sqvi["linhas"],
            "Est. Parq.": stock["linhas"],
            "LRR": _data_count(lrr_values),
            "Din": value_book["Din"].max_row,
            "Grafico": "gráficos e segmentadores, sem tabela operacional",
        },
        "base": base,
        "estoque": stock,
        "sqvi": sqvi,
        "cruzamentos": _cross(base, stock, sqvi, lrr_values),
    }


def _base_section(formula_rows, value_rows):
    items = []
    for excel_row, (formula, values) in enumerate(zip(formula_rows[1:], value_rows[1:]), start=2):
        if all(cell is None for cell in values):
            continue
        prepared = prepare_base_row(
            {
                "material": values[1],
                "center": values[3],
                "description": values[2],
                "plant": values[4],
                "lrr": values[5],
                "forecast": values[9],
                "status": values[13],
                "observation": values[16],
                "quantity": values[14],
                "delivered_on": values[15],
                "estoque_parque": values[6],
                "estoque_ok": values[7],
                "quant_po": values[8],
                "conc": values[0],
                "mes_ano": values[10],
                "coluna_mes": values[11],
                "coluna_ano": values[12],
            },
            excel_row,
        )
        referenced = referenced_base_row(formula[8])
        items.append(
            {
                "linha": excel_row,
                "material": prepared["material"],
                "centro": prepared["center"],
                "descricao": prepared["description"],
                "planta": text_from_code(values[4]),
                "lrr": prepared["lrr"],
                "previsao": prepared["forecast_on"],
                "previsao_original": prepared["forecast_original"],
                "status": prepared["status"],
                "observacao": prepared["observation"],
                "entrega": prepared["delivery"],
                "revisoes": prepared["revisoes"],
                "conc": text_from_code(values[0]),
                "po_cache": Decimal(str(values[8] or 0)),
                "estoque_cache": Decimal(str(values[6] or 0)),
                "formula_referencia": referenced,
                "formula_deslocada": referenced != excel_row,
            }
        )
    keys = [(item["material"], item["centro"]) for item in items]
    if len(keys) != len(set(keys)):
        raise SystemExit("A Base tem chave material + centro repetida.")
    origins = [{"delivery": item["entrega"]} for item in items]
    if repeated_origins(origins):
        raise SystemExit("A Base repetiu a origem de uma entrega.")
    return {"itens": items}


def _stock_section(rows):
    grouped = defaultdict(list)
    total = 0
    for excel_row, values in enumerate(rows[1:], start=2):
        if all(cell is None for cell in values):
            continue
        total += 1
        material = text_from_code(values[5])
        center = read_center(values[1])
        mapped = {
            "chave_planilha": text_from_code(values[0]),
            "nome": text_from_code(values[2]),
            "deposito": text_from_code(values[3]),
            "descricao_deposito": text_from_code(values[4]),
            "descricao": text_from_code(values[6]),
            "unidade": text_from_code(values[7]),
            "utilizacao_livre": values[9],
            "transito": values[10],
            "qualidade": values[11],
            "bloqueado": values[12],
        }
        key = (material, center, mapped["deposito"], text_from_code(values[8]))
        grouped[key].append({"linha": excel_row, "valores": mapped})
    classified = []
    for key, members in grouped.items():
        result = classify_position_group([member["valores"] for member in members])
        distinct = {
            field: sorted({_plain(member["valores"].get(field)) for member in members})
            for field in result["campos_divergentes"]
        }
        classified.append(
            {
                "material": key[0],
                "centro": key[1],
                "deposito": key[2],
                "lote": key[3],
                "linhas": [member["linha"] for member in members],
                "descricoes": sorted({member["valores"]["descricao"] for member in members}),
                "chaves": sorted({member["valores"]["chave_planilha"] for member in members}),
                "valores_divergentes": distinct,
                **result,
            }
        )
    classified.sort(key=lambda item: (item["situacao"], item["material"], item["centro"], item["deposito"], item["lote"]))
    return {"linhas": total, "grupos": classified}


def _sqvi_section(rows):
    movements = []
    seen = set()
    for excel_row, values in enumerate(rows[1:], start=2):
        if all(cell is None for cell in values):
            continue
        key = movement_key(values[9], values[10])
        if key in seen:
            raise SystemExit(f"Documento e item repetidos na linha {excel_row}.")
        seen.add(key)
        movements.append(
            {
                "linha": excel_row,
                "chave": key,
                "conc": text_from_code(values[1]),
                "tipo": text_from_code(values[7]),
                "centro": text_from_code(values[17]),
                "origem": text_from_code(values[13]),
                "material": text_from_code(values[19]),
                "pedido": Decimal(str(values[21])),
                "confirmada": Decimal(str(values[36])),
                "fornecimento": text_from_code(values[54]),
                "remessa": text_from_code(values[55]),
                "data_conf": values[39],
                "fim_em": values[44],
            }
        )
    return {"linhas": len(movements), "movimentos": movements}


def _cross(base, stock, sqvi, lrr_rows):
    points = {}
    for values in lrr_rows[1:]:
        if all(cell is None for cell in values):
            continue
        points[(text_from_code(values[1]).strip(), text_from_code(values[3]).strip())] = values[8]
    sheet_keys = {key for group in stock["grupos"] for key in group["chaves"]}
    descriptions = defaultdict(set)
    for group in stock["grupos"]:
        for text in group["descricoes"]:
            descriptions[(group["material"].strip(), group["centro"])].add(text)
    sums = defaultdict(lambda: Decimal("0"))
    counts = Counter()
    for movement in sqvi["movimentos"]:
        sums[movement["conc"]] += movement["pedido"]
        counts[movement["conc"]] += 1
    by_row = {item["linha"]: item for item in base["itens"]}
    relations = Counter()
    for item in base["itens"]:
        referenced = by_row.get(item["formula_referencia"])
        if referenced is None:
            raise SystemExit(
                f"A fórmula da linha {item['linha']} aponta para fora dos itens."
            )
        relation = po_cache_relation(
            item["po_cache"],
            sums[item["conc"]],
            sums[referenced["conc"]],
        )
        relations[relation] += 1
        item["soma_propria"] = sums[item["conc"]]
        item["relacao_po"] = relation
    lrr_equal = lrr_missing = lrr_diff = 0
    lrr_examples = []
    for item in base["itens"]:
        key = (item["material"].strip(), item["centro"])
        if key not in points:
            lrr_missing += 1
            continue
        point = points[key]
        if item["lrr"] is None and point in (None, ""):
            continue
        if item["lrr"] is None:
            lrr_examples.append((item["material"], item["centro"], point))
            continue
        if Decimal(str(point)) == item["lrr"]:
            lrr_equal += 1
        else:
            lrr_diff += 1
    absent = []
    key_failures = []
    description_conflicts = []
    for item in base["itens"]:
        key = (item["material"].strip(), item["centro"])
        groups = [
            group
            for group in stock["grupos"]
            if group["material"].strip() == key[0] and group["centro"] == key[1]
        ]
        if not groups:
            elsewhere = sorted(
                {
                    group["centro"]
                    for group in stock["grupos"]
                    if group["material"].strip() == key[0]
                }
            )
            absent.append((item["material"], item["centro"], elsewhere))
        if item["conc"] not in sheet_keys and groups:
            key_failures.append((item["material"], item["centro"], item["conc"]))
        other_texts = descriptions.get(key, ())
        if descriptions_for_review(item["descricao"], other_texts):
            description_conflicts.append(
                (item["material"], item["centro"], item["descricao"], tuple(sorted(other_texts)))
            )
    return {
        "relacoes_po": relations,
        "lrr_igual": lrr_equal,
        "lrr_vazio_com_ponto": lrr_examples,
        "lrr_ausente": lrr_missing,
        "lrr_diverge": lrr_diff,
        "sem_posicao": absent,
        "falha_de_chave": key_failures,
        "descricoes": description_conflicts,
        "movimentos_na_base": sum(1 for item in base["itens"] if counts[item["conc"]]),
        "linhas_sqvi_na_base": sum(counts[item["conc"]] for item in base["itens"]),
        "itens_com_varias_movimentacoes": sum(
            1 for item in base["itens"] if counts[item["conc"]] > 1
        ),
        "max_movimentacoes": max((counts[item["conc"]] for item in base["itens"]), default=0),
        "saldo_ausente": missing_stock(),
    }


def _print_report(report):
    print("Pré-validação da planilha aprovada")
    print(f"sha256 {report['arquivo']['sha256']}")
    print(f"bytes {report['arquivo']['bytes']}")
    print(f"planilha inalterada {report['arquivo']['sha256'] == report['arquivo_depois']['sha256']}")
    print("banco não acessado")
    print("Abas")
    for name, count in report["abas"].items():
        print(f"  {name}: {count}")
    items = report["base"]["itens"]
    print(f"Itens {len(items)}")
    print(f"LRR vazio {sum(1 for item in items if item['lrr'] is None)}")
    print(f"Previsão em data {sum(1 for item in items if item['previsao'] is not None)}")
    print(f"Previsão textual {sum(1 for item in items if item['previsao_original'])}")
    deliveries = [item["entrega"] for item in items if item["entrega"] is not None]
    print(f"Entregas {len(deliveries)}")
    print(f"Entregas negativas {sum(1 for delivery in deliveries if delivery['requires_review'])}")
    print(f"Entregas sem data {sum(1 for delivery in deliveries if delivery['delivered_on'] is None)}")
    equal_delivery = sum(
        1
        for item in items
        if item["entrega"] is not None and item["entrega"]["quantity"] == item["soma_propria"]
    )
    print(f"Entregas iguais à soma documental da própria chave {equal_delivery} de {len(deliveries)}")
    print(f"Fórmulas de Quant. PO deslocadas {sum(1 for item in items if item['formula_deslocada'])}")
    print(
        "Materiais com espaço",
        [item["material"] for item in items if item["material"] != item["material"].strip()],
    )
    plants = defaultdict(set)
    for item in items:
        plants[item["centro"]].add(item["planta"])
    print("Plantas", {center: sorted(names) for center, names in sorted(plants.items())})
    ties = [
        item
        for item in items
        if item["relacao_po"] == "empate" and item["po_cache"] == 0 and item["soma_propria"] == 0
    ]
    print("Relação do cache", dict(report["cruzamentos"]["relacoes_po"]))
    print(f"Empates em zero {len(ties)}")
    groups = report["estoque"]["grupos"]
    print(f"Grupos de estoque {len(groups)}")
    for situation in ("unica", "copia_identica", "conflito"):
        selected = [group for group in groups if group["situacao"] == situation]
        print(f"  {situation}: {len(selected)}")
    identical = [group for group in groups if group["situacao"] == "copia_identica"]
    for field, label in (
        ("utilizacao_livre", "utilização livre"),
        ("transito", "trânsito"),
        ("qualidade", "qualidade"),
        ("bloqueado", "bloqueado"),
    ):
        extra = sum(
            (
                group["efeito_se_somar"][field]
                - group["efeito_se_somar"][field] / group["ocorrencias"]
            )
            for group in identical
        )
        print(f"Aumento de {label} se as cópias idênticas fossem somadas {extra}")
    print("Cópias idênticas")
    for group in groups:
        if group["situacao"] == "copia_identica":
            print(_group_line(group))
    print("Conflitos")
    for group in groups:
        if group["situacao"] == "conflito":
            print(_group_line(group))
    cross = report["cruzamentos"]
    print(f"LRR igual ao ponto {cross['lrr_igual']}")
    print(f"LRR vazio com ponto {cross['lrr_vazio_com_ponto']}")
    print(f"LRR ausente {cross['lrr_ausente']}")
    print(f"Itens sem posição {len(cross['sem_posicao'])}")
    for material, center, elsewhere in cross["sem_posicao"]:
        print(f"  sem posição {material} {center} outros={elsewhere or '-'}")
    print(f"Saldo desses itens {cross['saldo_ausente']}")
    print(f"Falhas de chave {cross['falha_de_chave']}")
    print(f"Descrições para revisão {len(cross['descricoes'])}")
    for material, center, base_text, others in cross["descricoes"]:
        print(f"  descrição {material} {center} base={base_text} estoque={others}")
    print(f"Itens com movimentação {cross['movimentos_na_base']}")
    print(f"Linhas SQVI nesses itens {cross['linhas_sqvi_na_base']}")
    print(f"Itens com mais de uma movimentação {cross['itens_com_varias_movimentacoes']}")
    print(f"Máximo de movimentações {cross['max_movimentacoes']}")
    _print_sqvi(report["sqvi"])
    print("Nenhum registro de negócio foi proposto para gravação.")


def _print_sqvi(sqvi):
    types = Counter(movement["tipo"] for movement in sqvi["movimentos"])
    print("Centros de destino", dict(Counter(movement["centro"] for movement in sqvi["movimentos"])))
    print("Centros de origem", dict(Counter(movement["origem"] for movement in sqvi["movimentos"])))
    print("Tipos", dict(types))
    documents = {movement["chave"][0] for movement in sqvi["movimentos"]}
    print(f"Documentos {len(documents)}")
    print(f"Documentos com item {len({movement['chave'] for movement in sqvi['movimentos']})}")
    fornecimento = Counter(
        "#N/A" if movement["fornecimento"] == "#N/A" else "numero"
        for movement in sqvi["movimentos"]
    )
    print("Fornecimento", dict(fornecimento))
    confirmed = Counter()
    for movement in sqvi["movimentos"]:
        if movement["confirmada"] == 0:
            confirmed["zero"] += 1
        elif movement["confirmada"] == movement["pedido"]:
            confirmed["igual"] += 1
        elif movement["confirmada"] < movement["pedido"]:
            confirmed["menor"] += 1
        else:
            confirmed["maior"] += 1
    print("Quantidade confirmada", dict(confirmed))
    sentinels = [
        movement["chave"][0]
        for movement in sqvi["movimentos"]
        if excel_serial_calendar(movement["data_conf"]) == date(9999, 12, 31)
    ]
    print(f"DataConf em 31 dez 9999 {sentinels}")
    serials = sum(1 for movement in sqvi["movimentos"] if isinstance(movement["fim_em"], int))
    print(f"Data fim da EM em serial {serials}")


def _group_line(group):
    quantities = group["efeito_se_somar"]
    return (
        f"{group['situacao']} {group['material']} {group['centro']} "
        f"dep={group['deposito'] or '(vazio)'} lote={group['lote'] or '(vazio)'} "
        f"linhas={','.join(str(number) for number in group['linhas'])} "
        f"campos={','.join(group['campos_divergentes']) or '-'} "
        f"valores={group['valores_divergentes'] or '-'} "
        f"soma_se_juntar={{{', '.join(f'{name}: {_plain(value)}' for name, value in quantities.items())}}}"
    )


def _plain(value):
    if isinstance(value, Decimal):
        return format(value, "f")
    return value


def _rows(worksheet):
    return [tuple(cell.value for cell in row) for row in worksheet.iter_rows()]


def _require_headers(actual, expected, sheet):
    if tuple(actual[: len(expected)]) != tuple(expected):
        raise SystemExit(f"Cabeçalhos de {sheet} diferentes do esperado.")


def _data_count(rows):
    return sum(1 for row in rows[1:] if any(cell is not None for cell in row))


def _file_stamp(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    stat = path.stat()
    return {
        "bytes": stat.st_size,
        "sha256": digest.hexdigest(),
        "mtime": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
    }


def _configure_output():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")


if __name__ == "__main__":
    main()
