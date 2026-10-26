from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from criticals.models import CriticalItem, LinkedSheetRow, StockPosition
from criticals.source_integration import (
    SHEET_SPECS,
    STOCK_SOURCE,
    classify_existing_rows,
    read_complementary_sheets,
)


class ComplementarySheetReaderTests(TestCase):
    def setUp(self):
        self.item = CriticalItem.objects.create(
            material="MAT-1",
            center="B825",
            description="Synthetic item",
            status="Pendente",
        )
        self.temp_dir = TemporaryDirectory()
        self.workbook_path = Path(self.temp_dir.name) / "synthetic.xlsx"
        _write_workbook(self.workbook_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_exact_matches_preserve_space_negative_decimal_and_formula(self):
        prepared = read_complementary_sheets(
            self.workbook_path,
            {(self.item.material, self.item.center): self.item},
            require_official=False,
        )

        stock = prepared.sheets["Est. Parq."]
        self.assertEqual(len(stock.matched_rows), 1)
        self.assertEqual(stock.excluded_rows, 1)
        stock_row = stock.matched_rows[0]
        self.assertEqual(stock_row.source_row, 2)
        self.assertEqual(stock_row.stock_fields["unrestricted"], Decimal("1.25"))
        self.assertEqual(stock_row.stock_fields["in_transit"], Decimal("-2"))
        self.assertEqual(stock_row.stock_fields["in_quality"], Decimal("0"))
        self.assertEqual(stock_row.stock_fields["unit"], "PC")
        self.assertIn("=CONCAT(F2,B2)", stock_row.raw_values[0]["formula"])
        self.assertEqual(stock_row.material, "MAT-1")

    def test_sqvi_and_lrr_are_linked_by_exact_material_and_center(self):
        prepared = read_complementary_sheets(
            self.workbook_path,
            {(self.item.material, self.item.center): self.item},
            require_official=False,
        )

        self.assertEqual(len(prepared.sheets["SQVI"].matched_rows), 1)
        self.assertEqual(len(prepared.sheets["LRR"].matched_rows), 1)
        sqvi = prepared.sheets["SQVI"].matched_rows[0]
        self.assertEqual(sqvi.source_row, 2)
        self.assertEqual(sqvi.material, "MAT-1")
        self.assertEqual(sqvi.center, "B825")
        self.assertEqual(sqvi.raw_values[10]["header"], "Item")
        self.assertEqual(sqvi.raw_values[22]["header"], "UMP")

    def test_reimport_classification_never_updates_known_rows(self):
        prepared = read_complementary_sheets(
            self.workbook_path,
            {(self.item.material, self.item.center): self.item},
            require_official=False,
        )
        stock_row = prepared.sheets["Est. Parq."].matched_rows[0]
        sqvi_row = prepared.sheets["SQVI"].matched_rows[0]
        lrr_row = prepared.sheets["LRR"].matched_rows[0]
        existing_stock = [
            SimpleNamespace(
                source_row=stock_row.source_row,
                content_sha256=stock_row.content_sha256,
            )
        ]
        existing_linked = [
            SimpleNamespace(
                sheet=row.sheet,
                source_row=row.source_row,
                content_sha256=row.content_sha256,
            )
            for row in (
                _as_linked("SQVI", sqvi_row),
                _as_linked("LRR", lrr_row),
            )
        ]

        classified, conflicts = classify_existing_rows(
            prepared,
            existing_stock,
            existing_linked,
        )

        self.assertFalse(conflicts)
        self.assertEqual(
            [state for _row, state in classified["Est. Parq."]], ["igual"]
        )
        self.assertEqual([state for _row, state in classified["SQVI"]], ["igual"])
        self.assertEqual([state for _row, state in classified["LRR"]], ["igual"])

    def test_changed_content_is_a_conflict_not_an_update(self):
        prepared = read_complementary_sheets(
            self.workbook_path,
            {(self.item.material, self.item.center): self.item},
            require_official=False,
        )
        row = prepared.sheets["SQVI"].matched_rows[0]
        existing = [
            SimpleNamespace(
                sheet="SQVI",
                source_row=row.source_row,
                content_sha256="different-content",
            )
        ]

        _classified, conflicts = classify_existing_rows(prepared, [], existing)

        self.assertEqual(conflicts, [("SQVI", row.source_row)])

    def test_header_drift_is_refused(self):
        workbook = Workbook()
        stock = workbook.active
        stock.title = "Est. Parq."
        for index in range(13):
            stock.cell(row=1, column=index + 1, value=f"col{index}")
        for sheet_name in ("SQVI", "LRR"):
            sheet = workbook.create_sheet(sheet_name)
            for index in range(SHEET_SPECS[sheet_name]["column_count"]):
                sheet.cell(row=1, column=index + 1, value=f"col{index}")
        workbook.save(self.workbook_path)
        workbook.close()

        with self.assertRaisesRegex(ValueError, "Cabeçalho inesperado"):
            read_complementary_sheets(
                self.workbook_path,
                {(self.item.material, self.item.center): self.item},
                require_official=False,
            )


class LinkedSheetViewTests(TestCase):
    def setUp(self):
        self.item = CriticalItem.objects.create(
            material="MAT-1",
            center="B825",
            description="Synthetic item",
            status="Pendente",
        )
        self.user = get_user_model().objects.create_user(
            username="source-reader",
            password="test-password-only",
        )
        permission = Permission.objects.get(
            content_type__app_label="criticals",
            codename="view_criticalitem",
        )
        self.user.user_permissions.add(permission)
        self.client.force_login(self.user)

    def test_views_require_the_existing_item_read_permission(self):
        response = self.client.get(reverse("critical-sqvi-dashboard"))
        self.assertEqual(response.status_code, 200)
        self.user.user_permissions.clear()
        response = self.client.get(reverse("critical-lrr-dashboard"))
        self.assertEqual(response.status_code, 403)

    def test_source_rows_are_read_only_and_filter_to_linked_items(self):
        LinkedSheetRow.objects.create(
            item=self.item,
            sheet=LinkedSheetRow.Sheet.SQVI,
            source_row=2,
            material=self.item.material,
            center=self.item.center,
            raw_values=[
                {
                    "column": "J",
                    "header": "Doc.compra",
                    "value": {"type": "text", "value": "DOC-1"},
                }
            ],
            source_file_sha256="a" * 64,
            content_sha256="b" * 64,
        )

        response = self.client.get(
            reverse("critical-sqvi-dashboard"),
            {"material": "MAT-1", "centro": "B825"},
        )

        self.assertContains(response, "DOC-1")
        self.assertContains(response, "Correspondência exata")
        self.assertNotContains(response, "alterar status")

    def test_source_row_identity_is_unique_per_sheet(self):
        values = {
            "item": self.item,
            "sheet": LinkedSheetRow.Sheet.LRR,
            "source_row": 2,
            "material": self.item.material,
            "center": self.item.center,
            "raw_values": [],
            "source_file_sha256": "a" * 64,
            "content_sha256": "b" * 64,
        }
        LinkedSheetRow.objects.create(**values)

        with self.assertRaises(IntegrityError), transaction.atomic():
            LinkedSheetRow.objects.create(**values)


def _as_linked(sheet, row):
    return SimpleNamespace(
        sheet=sheet,
        source_row=row.source_row,
        content_sha256=row.content_sha256,
    )


def _write_workbook(path):
    workbook = Workbook()
    stock = workbook.active
    stock.title = "Est. Parq."
    stock_headers = [None] + [None] * 12
    for index, name in SHEET_SPECS["Est. Parq."]["signature"].items():
        stock_headers[index] = name
    stock.append(stock_headers)
    stock.append(
        [
            "=CONCAT(F2,B2)",
            "B825",
            "Centro de teste",
            "DEP1",
            "Depósito de teste",
            "MAT-1",
            "Descrição de teste",
            "PC",
            "LOTE1",
            1.25,
            -2,
            0,
            0,
        ]
    )
    stock.append(
        [
            "=CONCAT(F3,B3)",
            "B825",
            "Centro",
            "DEP1",
            "Depósito",
            "MAT-1 ",
            "Descrição",
            "PC",
            "L1",
            1,
            0,
            0,
            0,
        ]
    )

    sqvi = workbook.create_sheet("SQVI")
    sqvi_headers = [f"Campo {index + 1}" for index in range(61)]
    for index, name in SHEET_SPECS["SQVI"]["signature"].items():
        sqvi_headers[index] = name
    sqvi.append(sqvi_headers)
    sqvi_values = [None] * 61
    sqvi_values[0] = "=I2&J2"
    sqvi_values[1] = "=_xlfn.CONCAT(T2,R2)"
    sqvi_values[7] = "UB"
    sqvi_values[9] = "DOC-1"
    sqvi_values[10] = "10"
    sqvi_values[11] = "Fornecedor teste"
    sqvi_values[17] = "B825"
    sqvi_values[19] = "MAT-1"
    sqvi_values[21] = 2
    sqvi_values[22] = "PC"
    sqvi.append(sqvi_values)
    other_center = list(sqvi_values)
    other_center[17] = "B826"
    sqvi.append(other_center)

    lrr = workbook.create_sheet("LRR")
    lrr_headers = [f"Campo {index + 1}" for index in range(48)]
    for index, name in SHEET_SPECS["LRR"]["signature"].items():
        lrr_headers[index] = name
    lrr.append(lrr_headers)
    lrr_values = [None] * 48
    lrr_values[0] = "=CONCAT(B2,D2)"
    lrr_values[1] = "MAT-1"
    lrr_values[3] = "B825"
    lrr_values[5] = "ZY"
    lrr_values[6] = "ZRP"
    lrr_values[7] = "ZS20"
    lrr_values[8] = 1
    lrr.append(lrr_values)
    workbook.save(path)
    workbook.close()
