"""Pré-validação e carga futura da aba Base.

A carga real fica recusada sem autorização explícita e sem o banco local
confirmado. Esta etapa não importa a planilha operacional.
"""

from django.conf import settings
from django.db import connection, transaction

from criticals.models import CriticalItem, Delivery
from criticals.preparation import (
    BASE_SOURCE,
    ITEM_COMPARE_FIELDS,
    PreparationError,
    REIMPORT_POLICY,
    normalized_collisions,
    plan_reimport,
    prepare_base_row,
    repeated_origins,
)

BASE_HEADERS = (
    "Conc",
    "GP",
    "Descrição",
    "Parque",
    "Planta",
    "LRR",
    "Estoque Parque",
    "Estoque Ok/Nok",
    "Quant. PO ",
    "Previsão no Parque ",
    "Mes/Ano",
    "Coluna4",
    "Coluna3",
    "Status",
    "Quant. Entregue",
    "Data entrega ",
    "Obs",
)


class LoadRefused(PreparationError):
    """A carga real não foi autorizada ou o banco não é o destino local."""


def confirm_local_database():
    database = settings.DATABASES["default"]
    engine = database.get("ENGINE", "")
    name = database.get("NAME", "")
    user = database.get("USER", "")
    host = database.get("HOST", "")
    port = str(database.get("PORT", ""))
    if (
        engine != "django.db.backends.postgresql"
        or name != "apcluster"
        or user != "apcluster_app"
        or host not in ("localhost", "127.0.0.1")
        or port != "5432"
    ):
        raise LoadRefused("Banco de destino não confirmado. Operação recusada.")
    with connection.cursor() as cursor:
        cursor.execute(
            "SELECT current_database(), inet_server_addr()::text, inet_server_port()"
        )
        current_name, address, current_port = cursor.fetchone()
    if current_name != "apcluster" or current_port != 5432:
        raise LoadRefused("A conexão ativa não é o banco local apcluster.")
    address = (address or "").split("/")[0]
    if address not in ("127.0.0.1", "::1"):
        raise LoadRefused("O PostgreSQL não está escutando só em localhost.")
    return {"name": current_name, "address": address, "port": current_port}


def read_base_rows(path):
    """Lê somente a aba Base. Não grava o arquivo."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Base" not in workbook.sheetnames:
            raise PreparationError("A aba Base não está no arquivo.")
        worksheet = workbook["Base"]
        rows = [tuple(cell.value for cell in row) for row in worksheet.iter_rows()]
    finally:
        workbook.close()
    if tuple(rows[0][: len(BASE_HEADERS)]) != BASE_HEADERS:
        raise PreparationError("Os cabeçalhos da aba Base diferem do esperado.")
    prepared = []
    errors = []
    for excel_row, values in enumerate(rows[1:], start=2):
        if all(cell is None for cell in values):
            continue
        try:
            prepared.append(
                prepare_base_row(_cells(values), excel_row, source=BASE_SOURCE)
            )
        except PreparationError as exc:
            errors.append({"linha": excel_row, "motivo": str(exc)})
    return prepared, errors


def preview_against_database(prepared_rows):
    """Classifica sem gravar. Entregas manuais entram só na contagem preservada."""
    items = {
        (item.material, item.center): _item_snapshot(item)
        for item in CriticalItem.objects.all()
    }
    deliveries = {}
    manual = 0
    for delivery in Delivery.objects.select_related("item"):
        if delivery.source_row is None:
            manual += 1
            continue
        deliveries[(delivery.source, delivery.source_row)] = _delivery_snapshot(delivery)
    return classify_prepared(prepared_rows, items, deliveries, manual)


def classify_prepared(prepared_rows, items, deliveries, manual_count):
    if REIMPORT_POLICY["atualiza_campos"]:
        raise LoadRefused("A política confirmada não atualiza campos.")
    classified = []
    for prepared in prepared_rows:
        key = (prepared["material"], prepared["center"])
        delivery = prepared.get("delivery")
        existing_delivery = None
        if delivery is not None:
            existing_delivery = deliveries.get((delivery["source"], delivery["source_row"]))
        plan = plan_reimport(
            prepared,
            existing_item=items.get(key),
            existing_delivery=existing_delivery,
        )
        classified.append(plan)
    collisions = normalized_collisions(
        (row["material"], row["center"]) for row in prepared_rows
    )
    repeats = repeated_origins(prepared_rows)
    summary = _summary(classified, collisions, manual_count)
    summary["origens_repetidas"] = len(repeats)
    summary["centros"] = tuple(sorted({row["center"] for row in prepared_rows}))
    return {
        "linhas": classified,
        "colisoes": collisions,
        "origens_repetidas": repeats,
        "resumo": summary,
    }


def apply_prepared(prepared_rows, *, authorized):
    """Grava só com autorização. Não atualiza linha já importada nem apaga histórico."""
    if not authorized:
        raise LoadRefused(
            "A carga real está bloqueada. Use a pré-validação até haver autorização explícita."
        )
    if REIMPORT_POLICY["atualiza_campos"]:
        raise LoadRefused("A política confirmada não atualiza campos.")
    if repeated_origins(prepared_rows):
        raise LoadRefused("O lote contém a mesma origem mais de uma vez. Nada foi gravado.")
    created_items = 0
    created_deliveries = 0
    unchanged = 0
    conflicts = 0
    with transaction.atomic():
        for prepared in prepared_rows:
            key = (prepared["material"], prepared["center"])
            item = CriticalItem.objects.filter(material=key[0], center=key[1]).first()
            delivery = prepared.get("delivery")
            existing_delivery = None
            if delivery is not None:
                existing_delivery = Delivery.objects.filter(
                    source=delivery["source"],
                    source_row=delivery["source_row"],
                ).select_related("item").first()
            plan = plan_reimport(
                prepared,
                existing_item=_item_snapshot(item) if item else None,
                existing_delivery=(
                    _delivery_snapshot(existing_delivery) if existing_delivery else None
                ),
            )
            if plan["divergencias"]:
                conflicts += 1
                continue
            was_new = item is None
            if was_new:
                item = CriticalItem.objects.create(
                    material=prepared["material"],
                    center=prepared["center"],
                    description=prepared["description"],
                    lrr=prepared["lrr"],
                    forecast_on=prepared["forecast_on"],
                    forecast_original=prepared["forecast_original"],
                    status=prepared["status"],
                    observation=prepared["observation"],
                )
                created_items += 1
            created_delivery = False
            if delivery is not None and existing_delivery is None:
                Delivery.objects.create(
                    item=item,
                    quantity=delivery["quantity"],
                    delivered_on=delivery["delivered_on"],
                    source=delivery["source"],
                    source_row=delivery["source_row"],
                )
                created_deliveries += 1
                created_delivery = True
            if not was_new and not created_delivery:
                unchanged += 1
    return {
        "itens_criados": created_items,
        "entregas_criadas": created_deliveries,
        "inalterados": unchanged,
        "conflitos": conflicts,
        "atualizados": 0,
    }


def _cells(values):
    return {
        "material": values[1],
        "description": values[2],
        "center": values[3],
        "plant": values[4],
        "lrr": values[5],
        "estoque_parque": values[6],
        "estoque_ok": values[7],
        "quant_po": values[8],
        "forecast": values[9],
        "mes_ano": values[10],
        "coluna_mes": values[11],
        "coluna_ano": values[12],
        "status": values[13],
        "quantity": values[14],
        "delivered_on": values[15],
        "observation": values[16],
        "conc": values[0],
    }


def _item_snapshot(item):
    return {field: getattr(item, field) for field in ("material", "center", *ITEM_COMPARE_FIELDS)}


def _delivery_snapshot(delivery):
    return {
        "source": delivery.source,
        "source_row": delivery.source_row,
        "quantity": delivery.quantity,
        "delivered_on": delivery.delivered_on,
        "author": delivery.author_id,
        "material": delivery.item.material,
        "center": delivery.item.center,
    }


def _summary(classified, collisions, manual_count):
    counts = {
        "novo_item": 0,
        "item_existente": 0,
        "nova_entrega": 0,
        "entrega_conhecida": 0,
        "revisar": 0,
        "reconhecer": 0,
        "criar": 0,
    }
    for plan in classified:
        counts[plan["item"]] += 1
        if plan["entrega"]:
            counts[plan["entrega"]] += 1
        counts[plan["acao"]] += 1
    counts["colisoes"] = len(collisions)
    counts["entregas_manuais_preservadas"] = manual_count
    counts["atualizados"] = 0
    return counts
